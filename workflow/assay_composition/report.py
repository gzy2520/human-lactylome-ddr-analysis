#!/usr/bin/env python3
"""Source-backed Chinese report for the four within-assay composition targets."""
from pathlib import Path
import csv, json, hashlib
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'outputs/20260929_assay_composition'
def read(name):return list(csv.DictReader((OUT/name).open()))
m=read('comparison_metrics.csv');ext=read('external_metrics.csv');labs=read('labels/protein_labels_31.csv')
targets=['Proteome_DDR','Kla_DDR','Proteome_DNA_repair','Kla_DNA_repair']
names={'Proteome_DDR':'普通组 · DDR','Kla_DDR':'Kla 组 · DDR','Proteome_DNA_repair':'普通组 · DNA repair','Kla_DNA_repair':'Kla 组 · DNA repair'}
def metric(t,ev,w='Material'):return next(r for r in m if r['Target']==t and r['Evaluation']==ev and r['Weighting']==w)
def fmt(v,n=3):return f'{float(v):.{n}f}'
d=Document();d.core_properties.author='gzy';d.core_properties.last_modified_by='gzy';d.core_properties.comments='';sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(1.8)
style=d.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10.5);style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC');style.paragraph_format.line_spacing=1.15;style.paragraph_format.space_after=Pt(5)
for sn,size in [('Title',18),('Heading 1',14),('Heading 2',12)]:
 st=d.styles[sn];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Heiti SC')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def p(text):return d.add_paragraph(text)
def h(text):return d.add_heading(text,1)
def table(headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,txt in zip(t.rows[0].cells,headers):c.text=txt
 for row in rows:
  for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
 for i,row in enumerate(t.rows):
  prop=row._tr.get_or_add_trPr();prop.append(OxmlElement('w:cantSplit'))
  for cell in row.cells:
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3);para.paragraph_format.line_spacing=1.0
    for r in para.runs:r.font.size=Pt(9);r.bold=i==0
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','insideH']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');borders.append(e)
 t._tbl.tblPr.append(borders)
 return t
def fig(path,caption,width=17):
 q=d.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.keep_with_next=True;q.add_run().add_picture(str(OUT/path),width=Cm(width));p(caption)
def page():d.add_page_break()
d.add_heading('RNA 预测组学内 DDR 与 DNA repair 蛋白占比',0)
p('四目标分析报告｜图 1a 计数口径｜2026 年 9 月 29 日')
h('一、任务与结论')
p('本次分别预测普通蛋白质组、乳酸化蛋白质组各自内部的通路蛋白占比。四个目标独立训练，分子与分母均来自同一组学；不再计算普通组和 Kla 组之间的交集比值。输入仍为全部 18,332 个共有 Ensembl 基因。')
table(['目标','分子','分母'],[[names[t],('DDR' if t.endswith('_DDR') else 'DNA repair')+' 注释蛋白数','全部普通组检出蛋白数' if t.startswith('Proteome') else '全部 Kla 组检出蛋白数']for t in targets],[4,6,7])
p('所有比例均乘 100。这里的“占比”是检出蛋白种类的组成，不是通路活性、表达强度、修饰强度或位点占有率。')
rows=[]
for t in targets:
 a=metric(t,'Within-material test');b=metric(t,'Source held out');rows.append([names[t],fmt(a['MAE']),fmt(b['MAE']),fmt(b['R2']),fmt(b['BaselineMAE'])])
table(['目标','内部 MAE','跨来源 MAE','跨来源 R²','无 RNA MAE'],rows,[4.4,3,3.3,2.5,3.8])
p('表 1：材料等权、材料内供者等权；MAE 单位为百分点。无 RNA 基线使用各折训练材料标签的中位数。内部留出 374 条 RNA、19 种材料；跨来源为 14 个关联块留出，共 1,898 条 RNA、28 种材料。')
better=[names[t] for t in targets if float(metric(t,'Source held out')['MAE'])<float(metric(t,'Source held out')['BaselineMAE'])]
p('材料等权的跨来源 MAE 低于无 RNA 基线的目标：'+('、'.join(better) if better else '无')+'。这说明在该评价口径下存在预测增益；不能据此宣称对任意新来源或单个患者均可准确预测。四个目标全部列出，未按结果优劣筛除目标。')
p('外部 ESCC 肿瘤中，Kla-DDR 参考占比为 6.831%，预测均值为 6.750%，MAE 0.190 个百分点；Kla-DNA repair 分别为 4.797%、4.730% 和 0.159 个百分点，均优于无 RNA 基线。但癌旁两个 Kla 目标的 MAE 分别为 1.881 和 1.033，均未优于基线，外部效果不能推广到所有材料。')
p('本轮固定模型参数，不根据新目标或外部结果调参。既有内部划分有开发历史，因此称内部留出验证；外部 RNA 与质谱来自不同队列，仅能进行材料级一致性比较。')
page();h('二、数据、计数与建模方法')
p('冻结数据包括 31 个蛋白组材料组和 1,898 条 RNA 记录。31 组按 ReferenceKey 对应 28 个 RNA 材料；多个蛋白组组共用同一 RNA 材料时，沿用组占比的算术平均。每条 RNA 共享其材料标签，并没有逐个体配对的蛋白组占比。增加 RNA 数量不等于增加独立蛋白组标签。')
p('DDR 使用 GO:0006974，DNA repair 使用 GO:0006281，均沿 is_a / part_of 关系及冻结注释定义，分别涉及 2,660 和 2,135 个 UniProt BaseAccession。GO 集合只限定输出标签，不筛选输入基因。')
p('普通蛋白组保留图 1a 的 SourceProteinID 计数单位；当一个源 ID 映射到多个 UniProt 时，任一映射符合 GO 即判定命中，但源 ID 只计一次。Kla 组按去重 BaseAccession 计数。31 组 DDR 的分子与分母均已逐项核对，与冻结图 1a 一致。')
p('普通组肺和胎盘数据包含 Ensembl protein 到 UniProt 的非一一映射，源 ID 数分别为 14,022、12,395，合并为唯一 UniProt 后分别为 12,358、10,847。为保持图 1a 定义，本轮保留源计数单位；映射差异单列于 identifier_sensitivity.csv。这些来源的计数不宜解释为完全统一粒度的蛋白种类数量。')
rows=[]
for t in targets:
 vals=[float(r['ObservedPercent'])for r in labs if r['Target']==t];rows.append([names[t],f'{min(vals):.3f}–{max(vals):.3f}%'])
table(['目标','31 组参考占比范围'],rows,[7,10])
p('RNA 采用既有校正矩阵和全来源共有的 18,332 个 Ensembl ID，训练与外部均使用 log2(TPM+0.5)。不使用 Gene Symbol 作分析主键。随机森林回归采用 500 棵树、mtry=135、min.node.size=1、随机种子 25。训练抽样使材料等权、材料内供者等权、供者内记录等权。')
p('内部划分为 1,524 条训练、374 条留出。已知供者不跨划分；106 条供者身份未知的记录沿用伪 ID，因此不能声称所有供者身份均已核验。另运行 14 个来源关联块留出，来源块不等于 14 个独立实验室。全量模型用于外部预测，不能用其回代值代替内部验证结果。')
p('主要评价为 MAE、RMSE、R²及无 RNA 基线对照，同时提供逐记录与材料等权结果。误差在 ±1、±2、±5 个百分点内的比例只作描述；±1 是针对本轮较窄占比范围的辅助展示阈值，不是事先规定的成功标准，也不是分类正确率。')
page();h('三、内部留出表现')
fig(Path('figures/internal_test.png'),'图 1：四个模型的内部留出预测。每点为一条 RNA 记录；横轴是赋予该记录的共享材料标签。')
rows=[]
for t in targets:
 a=metric(t,'Within-material test','Record');b=metric(t,'Within-material test');rows.append([names[t],fmt(a['MAE']),fmt(a['R2']),fmt(b['MAE']),fmt(b['Within1'],1)+'%'])
table(['目标','记录 MAE','记录 R²','材料 MAE','材料 ±1 pp'],rows,[4.5,3,3,3,3.5])
p('内部结果回答：同一材料体系中换用留出的 RNA 记录，能否重现该材料的参考占比。它不能证明准确预测了这些受试者各自的蛋白组成。全量回代拟合与训练子集拟合另保存在逐记录结果表中，不作为独立验证证据。')
page();h('四、跨来源表现与基线')
fig(Path('figures/source_heldout.png'),'图 2：14 折来源关联块留出的预测；该折来源不参与该折训练。')
rows=[]
for t in targets:
 a=metric(t,'Source held out','Record');b=metric(t,'Source held out');rows.append([names[t],fmt(a['MAE']),fmt(a['R2']),fmt(b['MAE']),fmt(b['R2']),fmt(b['BaselineMAE'])])
table(['目标','记录 MAE','记录 R²','材料 MAE','材料 R²','基线 MAE'],rows,[4.2,2.5,2.5,2.7,2.5,2.6])
p('RNA 数量较多的材料在逐记录评价中权重较大；材料等权评价则提高小样本材料的相对权重，所以两种 R² 可能不同甚至符号相反。不能只选较好的一种口径，亦不能把较小的百分点误差直接当成对材料差异的充分解释。')
page();h('五、外部数据与结果')
p('外部质谱为 PXD063945：排除 NAC 后，普通组每种组织 25 列、Kla 组每种组织 22 列，分别在组学内部合并检出列表。RNA 为 GSE130078 的 23 位患者、46 份配对肿瘤及癌旁组织，样本 ID 与训练集无重叠。RNA 临床治疗前史尚未确认。两队列不是患者级配对，因此本节比较的是 RNA 预测与另一队列的材料参考比例。')
xlab=read('external_labels/external_labels.csv');rows=[]
for t in targets:
 for mat,cn in [('ESCC_untreated','肿瘤'),('ESCC_adjacent_untreated','癌旁')]:
  a=next(r for r in ext if r['Scope']=='newRNA' and r['Target']==t and r['Material']==mat);l=next(r for r in xlab if r['Target']==t and r['Material']==mat)
  rows.append([names[t],cn,l['Numerator']+'/'+l['Denominator'],fmt(a['Reference']),fmt(a['MeanPrediction']),fmt(a['MAE']),fmt(a['BaselineMAE'])])
table(['目标','组织','分子/分母','参考 %','预测 %','MAE','基线 MAE'],rows,[4,1.2,2.5,2.2,2.2,2.2,2.7])
p('表 4：每行 23 条 RNA，MAE 单位为百分点；基线为全部训练材料标签中位数。癌旁 Kla 的 DDR 和 DNA repair 分子为 42 和 28，包括一个未在普通组检出的蛋白；因为本轮分母属于 Kla 组，不再将其排除。每个目标仅有肿瘤、癌旁两个共享参考标签，不能把 46 条记录当作 46 个独立质谱真值，也不据此计算组内 R²。')
new=[r for r in ext if r['Scope']=='newRNA'];good=[r for r in new if float(r['MAE'])<float(r['BaselineMAE'])]
p(f'独立 RNA 外部比较共 8 个目标×组织组合，其中 {len(good)} 个组合的 MAE 低于无 RNA 基线。是否存在外部预测增益应以该比较为依据，不能仅以误差落在较宽阈值内判断。')
p('再以训练中既有 ESCC 材料占比直接作为肿瘤预测值，Kla-DDR 与 Kla-DNA repair 的 MAE 分别为 0.594 和 0.530 个百分点，仍高于本轮 RNA 模型的 0.190 和 0.159。这是同组织历史标签对照；癌旁没有对应训练材料，未强行套用肿瘤标签。')
p('另将训练中使用过的 95 条 TCGA ESCC RNA 与新的质谱材料标签比较，结果见 external_metrics.csv 中 reused_TCGA。该比较只考察更换质谱参考后的差异，不属于独立 RNA 外部测试。模型见过同类组织或另一队列的 ESCC RNA，不等于见过 GSE130078 的这些 RNA 记录。')
page();h('六、外部预测分布与结论')
fig(Path('figures/external_newRNA.png'),'图 3：新 RNA 队列的预测分布。红色菱形为新质谱材料参考比例，蓝点为 RNA 预测。')
p('本轮将输出改为组学内部组成比例，四个模型均可训练并输出数值。内部留出表现用于支持已知材料体系内的共享标签重现；跨来源评价与无 RNA 基线用于判断迁移预测增益；外部材料比较用于检查另一队列的参考标签是否一致。这三类证据回答不同问题。')
p('Kla 两个目标的跨来源材料等权 MAE 优于无 RNA 基线，但逐记录 R² 仍为负，故不能笼统称为通用高精度模型。外部队列只有两个材料标签且没有个体配对真值，需结合上一页所有组合的基线比较评价外部效果。当前没有验证基因对蛋白组成或乳酸化的因果调控，也没有验证患者级预测。')
p('四个目标都被完整保留。若选择某一模型作为汇报重点，应明确选择依据和评价口径，并同时展示其他目标及基线结果；不能将本轮较窄目标尺度带来的较小绝对误差解释为普遍提升。')
h('验收材料')
p('comparison_metrics.csv：四目标、四评价、两加权口径；all_predictions.csv.gz：逐 RNA 预测；external_metrics.csv：新 RNA 与复用 RNA 两种比较；labels/：31 组计数、28 材料标签、成员与映射敏感性；models/：模型、训练集特征重要性、分组指标及软件版本。')
p('附图：每个模型的 figures/top25.png / .pdf 展示全部 18,332 个输入基因中的训练子集重要性前 25 名。重要性为回归方差减少，不代表调控机制。')
# Remove inherited decorative paragraph borders; retain only functional table rules.
for element in [d.element,d.styles.element]:
 for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
dest=Path('/Users/gzy2520/Desktop/ML_图1a四目标_20260929.docx');d.save(dest)
(OUT/'report_manifest.json').write_text(json.dumps({'report':str(dest),'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'metrics':'comparison_metrics.csv','external_metrics':'external_metrics.csv'},ensure_ascii=False,indent=2))
print(dest)
