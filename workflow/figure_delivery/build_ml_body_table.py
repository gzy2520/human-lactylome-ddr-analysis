"""One-page editable manuscript table, using the frozen performance CSV."""
from pathlib import Path
import csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[2]
out = root / 'outputs/20261006_teacher_figures/main_figures'
rows = list(csv.DictReader((out / 'Table_1_ML_performance.csv').open()))
labels = {'Proteome_DDR': '普通蛋白组 DDR', 'Kla_DDR': '乳酸化蛋白组 DDR',
          'Proteome_DNA_repair': '普通蛋白组 DNA repair', 'Kla_DNA_repair': '乳酸化蛋白组 DNA repair'}
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.8)
for name in ['Normal', 'Title']:
    s = doc.styles[name]
    s.font.name, s.font.size, s.font.color.rgb = 'Arial Unicode MS', Pt(11 if name == 'Normal' else 15), RGBColor(0,0,0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
doc.styles['Title'].font.bold = True
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
doc.styles['Normal'].paragraph_format.space_after = Pt(8)
doc.add_paragraph('表1 RNA预测组学内DDR和DNA repair占比的模型性能', style='Title')
doc.add_paragraph('四个目标均在同一质谱组学内计数，分子为DDR或DNA repair集合对应的检出蛋白数，分母为该组学全部检出蛋白数。输入为各RNA来源共同覆盖的18,332个Ensembl基因；RNA记录共享材料层面的蛋白组参考占比。')
headers = ['预测目标', '内部测试\nMAE / R²', '整来源留出\nMAE / R²', '外部肿瘤\nMAE', '外部邻近组织\nMAE']
table = doc.add_table(rows=1, cols=5)
table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
widths = [1.85, 1.3, 1.3, 1.15, 1.3]
for col, w in zip(table.columns, widths): col.width = Inches(w)
for cell, text, w in zip(table.rows[0].cells, headers, widths):
    cell.text, cell.width = text, Inches(w)
for d in rows:
    cells = table.add_row().cells
    values = [labels[d['Target']],
              f"{float(d['InternalMAE']):.3f} / {float(d['InternalR2']):.3f}",
              f"{float(d['SourceMAE']):.3f} / {float(d['SourceR2']):.3f}",
              f"{float(d['ExternalTumorMAE']):.3f}", f"{float(d['ExternalAdjacentMAE']):.3f}"]
    for cell, text, w in zip(cells, values, widths):
        cell.text, cell.width = text, Inches(w)
for ri, row in enumerate(table.rows):
    for ci, cell in enumerate(row.cells):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ['top', 'left', 'bottom', 'right']:
            e = OxmlElement('w:' + edge)
            for k, v in {'val':'single','sz':'4','color':'D9D9D9'}.items(): e.set(qn('w:' + k),v)
            borders.append(e)
        pr.append(borders)
        margins = OxmlElement('w:tcMar')
        for side, size in [('top','110'),('bottom','110'),('left','90'),('right','90')]:
            e=OxmlElement('w:'+side);e.set(qn('w:w'),size);e.set(qn('w:type'),'dxa');margins.append(e)
        pr.append(margins)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = p.paragraph_format.space_before = Pt(0)
            for run in p.runs:
                run.font.size = Pt(10.5)
                run.font.bold = ri == 0
doc.add_paragraph()
notes = [
 '注：MAE单位为百分点，R²为决定系数。内部测试与整来源留出采用材料和供者平衡指标；外部列为对应材料内RNA记录的MAE。',
 '内部划分为1,524条训练RNA和374条测试RNA，按已知供者信息隔离。整来源留出为14个按来源复用关系划分的来源块轮流留出，共覆盖1,898条RNA记录；它不是额外新增的第三份独立数据集。',
 '外部材料比较使用GSE130078的46条RNA记录（23名供者的肿瘤与邻近组织），与PXD063945中筛选的手术治疗组质谱材料参照值比较。RNA与质谱并非患者配对，RNA治疗前史未完整确认，因此不能解释为患者级独立盲测。',
 '内部测试精度较高；整来源留出性能降低，提示跨来源外推仍有限。外部比较只涉及两个材料参照值，不能代表广泛外部泛化性能。',
]
for text in notes:
    doc.add_paragraph(text)
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.title = 'RNA预测组学内DDR和DNA repair占比的模型性能'
for style in doc.styles:
    if style.type == 1:
        ppr=style.element.find(qn('w:pPr'))
        if ppr is not None:
            for b in list(ppr.findall(qn('w:pBdr'))): ppr.remove(b)
for p in list(doc.paragraphs)+[p for row in table.rows for c in row.cells for p in c.paragraphs]:
    for run in p.runs:
        run.font.name='Arial Unicode MS'
        rf=run._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+key),'Arial Unicode MS')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']: rf.attrib.pop(qn('w:'+key),None)
doc.save(out / 'Table_1_ML_performance.docx')
